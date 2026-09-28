"""Протокол автоматизированной сверки по образцу Приложения 2 к ТЗ (образец обязателен, ТЗ 9.2) — PDF и DOCX.

    python -m ml.protocol <данные.json> <out.pdf|out.docx> [--format pdf|docx]

Данные собирает backend сервиса (service/backend/src/protocol-data.js): объект, версия и статус протокола, статус
загрузки стадий, сводка, строки по разделам, решения инспектора, карточки доказательств. Здесь — только вёрстка.

Разделы, как в образце: 1 — статус загрузки документов; 2 — сводная статистика с процентами; 3 — параметры, не
проверенные из-за отсутствия документов; 4 — критические и 5 — существенные нарушения (значения ПД / РД / ИД,
отклонение, решение инспектора); 6 — подозрения ИИ; 7 — резолютивная часть (рекомендации по подтверждённым нарушениям).
Дальше приложения по ТЗ 9.2, п. 4: карточки доказательств кандидатов (файл, SHA-256, страница, рамка, фрагмент и вырезка
страницы с подсветкой), проверенные отрицательные результаты, целостность комплекта.
"""
import argparse
import html
import io
import json
import sys

import numpy as np
import pymupdf

from ml import paths
from ml.render import context_region, open_pdf

RED = (220, 38, 38)
MAX_CROPS = 40
DECISION_COLOR = {"CONFIRMED_VIOLATION": "#b91c1c", "NEGATIVE_VERIFIED": "#4b5563", "CLARIFICATION_REQUIRED": "#b45309",
                  "PENDING": "#1d4ed8"}
STAGE_COLOR = {"UPLOADED": "#15803d", "PARTIAL": "#b45309", "MISSING": "#b91c1c"}


# ---------- Вырезки страниц с рамкой ----------

def crop_image(pdf_path, page_no, bbox, width=1100):
    """Область вокруг рамки доказательства с красной обводкой → JPEG (bytes) или None."""
    if not pdf_path or not bbox:
        return None
    try:
        doc = open_pdf(pdf_path)
    except Exception:
        return None
    try:
        page = doc[page_no - 1]
        region = context_region(bbox)
        area = pymupdf.Rect(region[0] * page.rect.width, region[1] * page.rect.height,
                            region[2] * page.rect.width, region[3] * page.rect.height)
        zoom = min(width / area.width, 200 / 72)
        pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), clip=area, alpha=False)
        img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n).copy()
        # рамка в координатах вырезки
        rw, rh = region[2] - region[0], region[3] - region[1]
        x0 = int((bbox[0] - region[0]) / rw * pix.width) - 4
        x1 = int((bbox[2] - region[0]) / rw * pix.width) + 4
        y0 = int((bbox[1] - region[1]) / rh * pix.height) - 4
        y1 = int((bbox[3] - region[1]) / rh * pix.height) + 4
        x0, y0 = max(0, x0), max(0, y0)
        x1, y1 = min(pix.width - 1, x1), min(pix.height - 1, y1)
        t = 3
        for (a, b, c, d) in ((y0, y0 + t, x0, x1), (y1 - t, y1, x0, x1), (y0, y1, x0, x0 + t), (y0, y1, x1 - t, x1)):
            img[max(0, a):max(0, b) + 1, max(0, c):max(0, d) + 1, :3] = RED
        out = pymupdf.Pixmap(pymupdf.csRGB, pix.width, pix.height, img.tobytes(), False)
        return out.tobytes("jpeg", jpg_quality=82)
    except Exception:
        return None
    finally:
        doc.close()


def evidence_crops(pl):
    crops = {}
    for card in pl["cards"]:
        for i, e in enumerate(card["evidence"]):
            if len(crops) >= MAX_CROPS:
                return crops
            png = crop_image(e.get("pdf_path"), e["page"], e.get("bbox_norm"))
            if png:
                crops[f"{card['finding_id']}_{i}"] = png
    return crops


# ---------- PDF (HTML → PyMuPDF Story) ----------

CSS = """
@font-face { font-family: F; src: url(f.ttf); }
@font-face { font-family: F; src: url(fb.ttf); font-weight: bold; }
body { font-family: F; font-size: 8.5pt; color: #111; }
h1 { font-size: 14pt; text-align: center; margin: 0 0 6pt 0; }
h2 { font-size: 11pt; margin: 12pt 0 4pt 0; }
h3 { font-size: 9.5pt; margin: 8pt 0 3pt 0; }
p { margin: 1.5pt 0; }
table { border-collapse: collapse; width: 100%; margin: 2pt 0 4pt 0; }
td, th { border: 0.5pt solid #9ca3af; padding: 2pt 3pt; vertical-align: top; }
th { background-color: #e5eaf1; font-weight: bold; text-align: left; }
.muted { color: #6b7280; }
.small { font-size: 7.5pt; }
.head td { border: none; padding: 1pt 4pt 1pt 0; }
.status { font-weight: bold; }
.card { margin: 6pt 0 2pt 0; }
"""


def e(v):
    return html.escape("—" if v is None or v == "" else str(v))


def table(headers, rows, widths=None):
    ws = widths or [None] * len(headers)
    th = "".join(f"<th{f' width={w!r}' if w else ''}>{e(h)}</th>" for h, w in zip(headers, ws))
    body = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>" for r in rows)
    return f"<table><tr>{th}</tr>{body}</table>"


def colored(text, color, bold=False):
    b = " font-weight: bold;" if bold else ""
    return f'<span style="color: {color};{b}">{e(text)}</span>'


def findings_table(items):
    rows = [[e(f["no"]), e(f["section"]), f"{e(f['name'])} ({e(f['code'])})", e(f["location"]), e(f["pd"]), e(f["rd"]), e(f["id"]),
             e(f["deviation"]),
             colored(f["decision"], DECISION_COLOR.get(f["decision_status"], "#111"), True)
             + (f"<br/><span class='small muted'>{e(f['comment'])}</span>" if f.get("comment") else "")]
            for f in items]
    return table(["№", "Раздел", "Параметр (код)", "Место", "ПД", "РД", "ИД", "Отклонение", "Решение инспектора"], rows,
                 ["3%", "5%", "17%", "9%", "11%", "11%", "8%", "20%", "16%"])


def html_document(pl, crops):
    o, p = pl["object"], pl["protocol"]
    parts = [f"<h1>{e(pl['title'])} № {e(pl['number'])}</h1>"]
    head = [("Объект", o["name"]), ("Адрес", o["address"]), ("Номер надзорного дела", o["case_number"]),
            ("Застройщик", o["developer"]), ("Подрядчик", o["contractor"]), ("Дата формирования", p["created_at"]),
            ("Версия протокола", p["version_text"]),
            ("Статус", p["headline"]), ("Тип проверки", f"{p['scenario_text']} ({p['scenario']})")]
    if p.get("finalized_at"):
        head.append(("Финализирован", f"{p['finalized_at']}, {p.get('finalized_by') or '—'}"))
    parts.append("<table class='head'>" + "".join(
        f"<tr><td width='22%'><b>{e(k)}:</b></td><td>{e(v) if k != 'Статус' else colored(v, '#b45309' if p['status'] != 'PROTOCOL_FINALIZED' else '#15803d', True)}</td></tr>"
        for k, v in head) + "</table>")
    parts.append(f"<p class='small muted'>Код объекта: {e(o['id'])}. {e(p['matrix_version'])}. Модель: {e(p['model_version'])}. "
                 f"Набор данных: {e(p['dataset_version'])}. SHA-256 реестра входных файлов: {e(p['input_manifest_hash'])}.</p>")
    parts.append("<p class='small muted'>Система не выносит вердикт: строки разделов 4–5 — кандидаты в нарушения с доказательствами; "
                 "нарушением запись становится только после подтверждения инспектором (статус CONFIRMED_VIOLATION).</p>")

    # 1
    parts.append("<h2>РАЗДЕЛ 1. СТАТУС ЗАГРУЗКИ ДОКУМЕНТОВ</h2>")
    parts.append(table(["Тип документа", "Статус", "Загружено файлов", "Ожидается", "Комментарий"], [
        [e(s["stage_short"]), colored(f"{s['status_text']} ({s['status']})", STAGE_COLOR[s["status"].split("_")[1]], True),
         e(s["files"]), "—", e(s["comment"])] for s in pl["stages"]], ["10%", "22%", "13%", "10%", "45%"]))
    parts.append("<p class='small muted'>«Ожидается»: реестра ожидаемых файлов в комплекте нет, поэтому «частично» определяется по "
                 "проверке комплектности — разделы ПД без комплектов РД, нечитаемые файлы.</p>")
    # 2
    parts.append("<h2>РАЗДЕЛ 2. СВОДНАЯ СТАТИСТИКА (С ПРОЦЕНТАМИ)</h2>")
    parts.append(table(["Показатель", "Количество", "% от общего"],
                       [[e(r["label"]), e(r["count"]), e(r["share"])] for r in pl["summary"]], ["62%", "18%", "20%"]))
    parts.append("<p class='small muted'>Строки «проверено», «не проверено», «ручная проверка», «не проверялось» — по параметрам "
                 "матрицы (у параметра бывает несколько мест: конструкции, помещения); кандидаты и решения — по доказательным "
                 "группам (параметр + место). Проценты — от числа параметров матрицы.</p>")
    # 3
    parts.append(f"<h2>РАЗДЕЛ 3. ПАРАМЕТРЫ, НЕ ПРОВЕРЕННЫЕ ИЗ-ЗА ОТСУТСТВИЯ ДОКУМЕНТОВ — {len(pl['missing'])}</h2>")
    parts.append(table(["№", "Код", "Раздел", "Параметр", "Отсутствующий документ"],
                       [[e(m["no"]), e(m["code"]), e(m["section"]), e(m["name"]), e(m["missing"])] for m in pl["missing"]],
                       ["5%", "10%", "10%", "40%", "35%"]) if pl["missing"] else "<p>Таких параметров нет.</p>")
    # 4, 5
    parts.append(f"<h2>РАЗДЕЛ 4. КРИТИЧЕСКИЕ НАРУШЕНИЯ — {len(pl['critical'])}</h2>")
    parts.append(findings_table(pl["critical"]) if pl["critical"] else "<p>Кандидатов в критические нарушения нет.</p>")
    parts.append(f"<h2>РАЗДЕЛ 5. СУЩЕСТВЕННЫЕ НАРУШЕНИЯ — {len(pl['significant'])}</h2>")
    parts.append(findings_table(pl["significant"]) if pl["significant"] else "<p>Кандидатов в существенные нарушения нет.</p>")
    # 6
    parts.append(f"<h2>РАЗДЕЛ 6. ПОДОЗРЕНИЯ ИИ — {len(pl['suspicions'])}</h2>")
    if pl["suspicions"]:
        parts.append("<p class='small muted'>Сравнение выполнено, но уверенности нет: значение найдено только на одной стадии, "
                     "сработала страховка или документы расходятся между собой. Требуют ручной проверки.</p>")
        parts.append(table(["№", "Метод", "Описание", "ПД", "РД", "ИД", "Решение инспектора", "Причина отклонения", "Комментарий ИИ"],
                           [[e(s["no"]), e(s["method"]), e(s["description"]), e(s["pd"]), e(s["rd"]), e(s["id"]), e(s["decision"]),
                             e(s["reason"]), f"<span class='small'>{e(s['ai_comment'])}</span>"] for s in pl["suspicions"]],
                           ["3%", "10%", "22%", "10%", "10%", "7%", "10%", "10%", "18%"]))
    else:
        parts.append("<p>Подозрений нет.</p>")
    # 7
    parts.append("<h2>РАЗДЕЛ 7. РЕЗОЛЮТИВНАЯ ЧАСТЬ</h2>")
    for title, key, col in (("7.1. По критическим нарушениям", "critical", "Вид работ"),
                            ("7.2. По существенным нарушениям", "significant", "Вид нарушения")):
        parts.append(f"<h3>{title}</h3>")
        items = pl["resolution"][key]
        parts.append(table(["№", col, "Конкретная рекомендация"], [[e(r["no"]), e(r["work"]), e(r["recommendation"])] for r in items],
                           ["5%", "35%", "60%"]) if items else "<p>Подтверждённых инспектором нарушений нет.</p>")
    parts.append("<p style='margin-top: 14pt;'>Инспектор: ______________________ / ______________________ /"
                 "&#160;&#160;&#160;&#160; Дата: ____________</p>")
    parts.append(f"<p class='small muted'>Сформировано системой «Инспектор ИИ» {e(pl['generated_at'])}"
                 f"{', пользователь ' + e(pl['generated_by']) if pl.get('generated_by') else ''}.</p>")
    return "\n".join(parts)


def html_tail(pl):
    """Приложения Б и В (карточки доказательств — приложение А — рисуются отдельно, по странице на кандидата)."""
    parts = [f"<h2>ПРИЛОЖЕНИЕ Б. ПРОВЕРЕННЫЕ ОТРИЦАТЕЛЬНЫЕ РЕЗУЛЬТАТЫ (NEGATIVE_VERIFIED) — {len(pl['negatives'])}</h2>"]
    parts.append(table(["№", "Код", "Параметр", "Место", "ПД", "РД", "ИД", "Доказательства"],
                       [[e(n["no"]), e(n["code"]), e(n["name"]), e(n["location"]), e(n["pd"]), e(n["rd"]), e(n["id"]),
                         f"<span class='small'>{e(n['evidence'])}</span>"] for n in pl["negatives"]],
                       ["4%", "8%", "20%", "14%", "12%", "12%", "10%", "20%"]) if pl["negatives"] else "<p>Нет.</p>")
    parts.append(f"<h2>ПРИЛОЖЕНИЕ В. ЦЕЛОСТНОСТЬ КОМПЛЕКТА ДОКУМЕНТОВ — {len(pl['integrity'])}</h2>")
    parts.append(table(["Важность", "Тип", "Находка"], [[e(x["severity"]), f"<span class='small'>{e(x['type'])}</span>", e(x["title"])]
                                                          for x in pl["integrity"]], ["10%", "25%", "65%"])
                 if pl["integrity"] else "<p>Дефектов комплекта не найдено.</p>")
    return "\n".join(parts)


PAGE = pymupdf.paper_rect("a4-l")


def _archive():
    arch = pymupdf.Archive()
    arch.add(open(paths.font_file(), "rb").read(), "f.ttf")
    arch.add(open(paths.font_file(bold=True), "rb").read(), "fb.ttf")
    return arch


def story_pdf(html_text):
    story = pymupdf.Story(html=html_text, user_css=CSS, archive=_archive())
    buf = io.BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    where = PAGE + (32, 30, -32, -34)
    more = 1
    while more:
        dev = writer.begin_page(PAGE)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    return pymupdf.open("pdf", buf.getvalue())


def card_fields(c):
    return [("Идентификатор группы", c["finding_id"]), ("Критичность / уровень риска", f"{c['criticality']} / {c['review_priority']}"),
            ("Ожидается (ПД)", c["expected"]), ("Фактически (РД)", c["actual_rd"]), ("Фактически (ИД)", c["actual_id"]),
            ("Триггер матрицы", c["trigger"]), ("Пояснение системы", c.get("note")),
            ("Решение инспектора", c["decision"] + (f": {c['reason']}" if c.get("reason") else "")),
            ("Комментарий инспектора", c.get("comment")),
            ("Кто и когда", f"{c['decided_by']}, {c['decided_at']}" if c.get("decided_by") else None),
            ("Комментарий ИИ", c.get("ai_comment"))]


def card_pages(doc, pl, crops):
    """Приложение А: по странице на кандидата — слева поля карточки и файлы, справа вырезки страниц с рамкой."""
    arch = _archive()
    total = len(pl["cards"])
    for n, c in enumerate(pl["cards"], 1):
        page = doc.new_page(width=PAGE.width, height=PAGE.height)
        head = (f"<p class='small muted'>ПРИЛОЖЕНИЕ А. КАРТОЧКИ ДОКАЗАТЕЛЬСТВ КАНДИДАТОВ · {n} из {total}</p>"
                f"<h3>{e(c['finding_no'])}. {e(c['name'])} ({e(c['code'])}) — {e(c['location'])}</h3>")
        page.insert_htmlbox(pymupdf.Rect(32, 26, PAGE.width - 32, 70), head, css=CSS, archive=arch)
        fields = "<table>" + "".join(f"<tr><td width='34%'><b>{e(k)}</b></td><td>{e(v)}</td></tr>" for k, v in card_fields(c) if v) + "</table>"
        page.insert_htmlbox(pymupdf.Rect(32, 72, 382, 352), fields, css=CSS, archive=arch, scale_low=0.5)
        rows = []
        for ev in c["evidence"]:
            box = ", ".join(f"{v:.3f}" for v in ev["bbox_norm"]) if ev.get("bbox_norm") else "—"
            rows.append(f"<tr><td><b>{e(ev['stage'])}</b></td><td>{e(ev['file_id'])}, стр. {e(ev['page'])}<br/>"
                        f"<span class='small'>{e(ev['file_name'])}</span><br/>"
                        f"<span class='small muted'>SHA-256 {e((ev.get('sha256') or '')[:24])}…<br/>рамка {box}</span></td>"
                        f"<td><span class='small'>{e(ev.get('fragment'))}</span></td></tr>")
        evt = ("<table><tr><th>Стадия</th><th>Файл и страница</th><th>Фрагмент в рамке</th></tr>" + "".join(rows) + "</table>")
        page.insert_htmlbox(pymupdf.Rect(32, 356, 382, PAGE.height - 34), evt, css=CSS, archive=arch, scale_low=0.4)
        # вырезки справа, одна под другой
        imgs = [(i, ev) for i, ev in enumerate(c["evidence"]) if f"{c['finding_id']}_{i}" in crops]
        if not imgs:
            page.insert_htmlbox(pymupdf.Rect(396, 72, PAGE.width - 32, 120), "<p class='muted'>Вырезки страниц недоступны (нет файла или рамки).</p>",
                                css=CSS, archive=arch)
            continue
        top, bottom, left, right = 72, PAGE.height - 34, 396, PAGE.width - 32
        slot = (bottom - top) / len(imgs)
        for k, (i, ev) in enumerate(imgs):
            y = top + k * slot
            page.insert_htmlbox(pymupdf.Rect(left, y, right, y + 14),
                                f"<p class='small'><b>{e(ev['stage'])}</b>: {e(ev['file_id'])}, стр. {e(ev['page'])} — фрагмент обведён красным</p>",
                                css=CSS, archive=arch)
            box = pymupdf.Rect(left, y + 15, right, y + slot - 6)
            page.insert_image(box, stream=crops[f"{c['finding_id']}_{i}"], keep_proportion=True)
            page.draw_rect(box, color=(0.8, 0.8, 0.8), width=0.4)


def build_pdf(pl, crops, out):
    doc = story_pdf(html_document(pl, crops))
    if pl["cards"]:
        card_pages(doc, pl, crops)
    tail = story_pdf(html_tail(pl))
    doc.insert_pdf(tail)
    tail.close()
    doc.set_metadata({"title": f"{pl['title']} № {pl['number']}", "author": "Инспектор ИИ", "subject": pl["object"]["name"]})
    # номер протокола и нумерация страниц внизу (один шрифт на документ: TextWriter встраивает его один раз)
    font = pymupdf.Font(fontfile=paths.font_file())
    for i, page in enumerate(doc, 1):
        tw = pymupdf.TextWriter(page.rect)
        tw.append((32, page.rect.height - 14), f"Протокол № {pl['number']} · версия {pl['protocol']['version_text']}",
                  font=font, fontsize=7)
        tw.append((page.rect.width - 90, page.rect.height - 14), f"стр. {i} из {doc.page_count}", font=font, fontsize=7)
        tw.write_text(page, color=(0.45, 0.45, 0.45))
    n = doc.page_count
    doc.save(out, garbage=4, deflate=True)          # garbage=4 схлопывает одинаковые шрифты insert_htmlbox
    doc.close()
    return n


# ---------- DOCX ----------

def build_docx(pl, crops, out):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.5))
    sec.top_margin = sec.bottom_margin = Cm(1.3)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(2)

    def rgb(hex_color):
        h = hex_color.lstrip("#")
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    def para(text="", bold=False, size=None, color=None, align=None, italic=False):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = rgb(color)
        if align:
            p.alignment = align
        return p

    def heading(text):
        p = para(text, bold=True, size=12)
        p.paragraph_format.space_before = Pt(10)

    def shade(cell, fill="E5EAF1"):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)

    def tbl(headers, rows, widths_cm=None, colors=None):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = ""
            run = c.paragraphs[0].add_run(h)
            run.bold = True
            run.font.size = Pt(9)
            shade(c)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                run = cells[i].paragraphs[0].add_run("—" if v is None or v == "" else str(v))
                run.font.size = Pt(9)
                if colors and colors[ri] and i in colors[ri]:
                    run.font.color.rgb = rgb(colors[ri][i])
                    run.bold = True
        if widths_cm:
            for row in t.rows:
                for i, w in enumerate(widths_cm):
                    row.cells[i].width = Cm(w)
        return t

    o, p = pl["object"], pl["protocol"]
    para(f"{pl['title']} № {pl['number']}", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    head = [("Объект", o["name"]), ("Адрес", o["address"]), ("Номер надзорного дела", o["case_number"]),
            ("Застройщик", o["developer"]), ("Подрядчик", o["contractor"]), ("Дата формирования", p["created_at"]),
            ("Версия протокола", p["version_text"]), ("Статус", p["headline"]),
            ("Тип проверки", f"{p['scenario_text']} ({p['scenario']})")]
    if p.get("finalized_at"):
        head.append(("Финализирован", f"{p['finalized_at']}, {p.get('finalized_by') or '—'}"))
    for k, v in head:
        q = doc.add_paragraph()
        q.add_run(f"{k}: ").bold = True
        r = q.add_run(v or "—")
        if k == "Статус":
            r.bold = True
            r.font.color.rgb = rgb("#15803d" if p["status"] == "PROTOCOL_FINALIZED" else "#b45309")
    para(f"Код объекта: {o['id']}. {p['matrix_version']}. Модель: {p['model_version']}. Набор данных: {p['dataset_version']}. "
         f"SHA-256 реестра входных файлов: {p['input_manifest_hash']}.", size=8, color="#6b7280")
    para("Система не выносит вердикт: строки разделов 4–5 — кандидаты в нарушения с доказательствами; нарушением запись "
         "становится только после подтверждения инспектором.", size=8, color="#6b7280", italic=True)

    heading("РАЗДЕЛ 1. СТАТУС ЗАГРУЗКИ ДОКУМЕНТОВ")
    tbl(["Тип документа", "Статус", "Загружено файлов", "Ожидается", "Комментарий"],
        [[s["stage_short"], f"{s['status_text']} ({s['status']})", s["files"], "—", s["comment"]] for s in pl["stages"]],
        [3, 5.5, 3, 2.5, 12.5], [{1: STAGE_COLOR[s["status"].split("_")[1]]} for s in pl["stages"]])
    heading("РАЗДЕЛ 2. СВОДНАЯ СТАТИСТИКА (С ПРОЦЕНТАМИ)")
    tbl(["Показатель", "Количество", "% от общего"], [[r["label"], r["count"], r["share"]] for r in pl["summary"]], [16, 4, 4])
    heading(f"РАЗДЕЛ 3. ПАРАМЕТРЫ, НЕ ПРОВЕРЕННЫЕ ИЗ-ЗА ОТСУТСТВИЯ ДОКУМЕНТОВ — {len(pl['missing'])}")
    if pl["missing"]:
        tbl(["№", "Код", "Раздел", "Параметр", "Отсутствующий документ"],
            [[m["no"], m["code"], m["section"], m["name"], m["missing"]] for m in pl["missing"]], [1.2, 2.3, 2.3, 11, 9.7])
    else:
        para("Таких параметров нет.")

    def findings(items):
        tbl(["№", "Раздел", "Параметр (код)", "Место", "ПД", "РД", "ИД", "Отклонение", "Решение инспектора"],
            [[f["no"], f["section"], f"{f['name']} ({f['code']})", f["location"], f["pd"], f["rd"], f["id"], f["deviation"],
              f["decision"] + (f"\n{f['comment']}" if f.get("comment") else "")] for f in items],
            [0.9, 1.4, 4.6, 2.4, 3, 3, 2.2, 5.5, 3.5],
            [{8: DECISION_COLOR.get(f["decision_status"], "#111111")} for f in items])

    heading(f"РАЗДЕЛ 4. КРИТИЧЕСКИЕ НАРУШЕНИЯ — {len(pl['critical'])}")
    findings(pl["critical"]) if pl["critical"] else para("Кандидатов в критические нарушения нет.")
    heading(f"РАЗДЕЛ 5. СУЩЕСТВЕННЫЕ НАРУШЕНИЯ — {len(pl['significant'])}")
    findings(pl["significant"]) if pl["significant"] else para("Кандидатов в существенные нарушения нет.")
    heading(f"РАЗДЕЛ 6. ПОДОЗРЕНИЯ ИИ — {len(pl['suspicions'])}")
    if pl["suspicions"]:
        tbl(["№", "Метод", "Описание", "ПД", "РД", "ИД", "Решение инспектора", "Причина отклонения", "Комментарий ИИ"],
            [[s["no"], s["method"], s["description"], s["pd"], s["rd"], s["id"], s["decision"], s["reason"], s["ai_comment"]]
             for s in pl["suspicions"]], [0.9, 2.6, 5.6, 2.6, 2.6, 1.8, 2.6, 2.6, 5.2])
    else:
        para("Подозрений нет.")
    heading("РАЗДЕЛ 7. РЕЗОЛЮТИВНАЯ ЧАСТЬ")
    for title, key, col in (("7.1. По критическим нарушениям", "critical", "Вид работ"),
                            ("7.2. По существенным нарушениям", "significant", "Вид нарушения")):
        para(title, bold=True)
        items = pl["resolution"][key]
        if items:
            tbl(["№", col, "Конкретная рекомендация"], [[r["no"], r["work"], r["recommendation"]] for r in items], [1.2, 9, 16.3])
        else:
            para("Подтверждённых инспектором нарушений нет.")
    para("")
    para("Инспектор: ______________________ / ______________________ /        Дата: ____________")
    para(f"Сформировано системой «Инспектор ИИ» {pl['generated_at']}"
         + (f", пользователь {pl['generated_by']}" if pl.get("generated_by") else "") + ".", size=8, color="#6b7280")

    doc.add_page_break()
    heading(f"ПРИЛОЖЕНИЕ А. КАРТОЧКИ ДОКАЗАТЕЛЬСТВ КАНДИДАТОВ — {len(pl['cards'])}")
    for c in pl["cards"]:
        para(f"{c['finding_no']}. {c['name']} ({c['code']}) — {c['location']}", bold=True, size=10.5)
        rows = [("Идентификатор группы", c["finding_id"]), ("Критичность / уровень риска", f"{c['criticality']} / {c['review_priority']}"),
                ("Ожидается (ПД)", c["expected"]), ("Фактически (РД)", c["actual_rd"]), ("Фактически (ИД)", c["actual_id"]),
                ("Триггер матрицы", c["trigger"]), ("Пояснение системы", c.get("note")),
                ("Решение инспектора", c["decision"] + (f": {c['reason']}" if c.get("reason") else "")),
                ("Комментарий инспектора", c.get("comment")),
                ("Кто и когда", f"{c['decided_by']}, {c['decided_at']}" if c.get("decided_by") else None),
                ("Комментарий ИИ", c.get("ai_comment"))]
        tbl(["Поле", "Значение"], [[k, v] for k, v in rows if v], [6, 20.5])
        tbl(["Стадия", "Файл", "SHA-256", "Стр.", "Рамка", "Фрагмент"],
            [[ev["stage"], f"{ev['file_id']} {ev['file_name']}", (ev.get("sha256") or "")[:16] + "…", ev["page"],
              ", ".join(f"{v:.3f}" for v in ev["bbox_norm"]) if ev.get("bbox_norm") else "—", ev.get("fragment")]
             for ev in c["evidence"]], [1.5, 7, 3.5, 1.2, 4, 9.3])
        for i, ev in enumerate(c["evidence"]):
            png = crops.get(f"{c['finding_id']}_{i}")
            if png:
                para(f"{ev['stage']}: {ev['file_id']}, стр. {ev['page']}", size=8.5, bold=True)
                doc.add_picture(io.BytesIO(png), width=Cm(15))
    heading(f"ПРИЛОЖЕНИЕ Б. ПРОВЕРЕННЫЕ ОТРИЦАТЕЛЬНЫЕ РЕЗУЛЬТАТЫ (NEGATIVE_VERIFIED) — {len(pl['negatives'])}")
    if pl["negatives"]:
        tbl(["№", "Код", "Параметр", "Место", "ПД", "РД", "ИД", "Доказательства"],
            [[n["no"], n["code"], n["name"], n["location"], n["pd"], n["rd"], n["id"], n["evidence"]] for n in pl["negatives"]],
            [1, 2, 5.5, 3.5, 3.2, 3.2, 2.6, 5.5])
    heading(f"ПРИЛОЖЕНИЕ В. ЦЕЛОСТНОСТЬ КОМПЛЕКТА ДОКУМЕНТОВ — {len(pl['integrity'])}")
    if pl["integrity"]:
        tbl(["Важность", "Тип", "Находка"], [[x["severity"], x["type"], x["title"]] for x in pl["integrity"]], [2.5, 7, 17])
    doc.core_properties.title = f"{pl['title']} № {pl['number']}"
    doc.core_properties.author = "Инспектор ИИ"
    doc.save(out)


def build(payload_path, out, fmt="pdf"):
    with open(payload_path, encoding="utf-8") as f:
        pl = json.load(f)
    crops = evidence_crops(pl)
    if fmt == "pdf":
        build_pdf(pl, crops, out)
    elif fmt == "docx":
        build_docx(pl, crops, out)
    else:
        raise ValueError(f"формат {fmt}: pdf или docx")
    return out


def main(argv):
    ap = argparse.ArgumentParser(description="Протокол по образцу Приложения 2 (PDF, DOCX)")
    ap.add_argument("payload")
    ap.add_argument("out")
    ap.add_argument("--format", choices=["pdf", "docx"])
    a = ap.parse_args(argv)
    fmt = a.format or ("docx" if a.out.lower().endswith(".docx") else "pdf")
    build(a.payload, a.out, fmt)
    print(a.out)


if __name__ == "__main__":
    main(sys.argv[1:])
